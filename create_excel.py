import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

csv_path = r"C:\Users\maddi\OneDrive\Desktop\DataI2I\Restaurant-Management-System\User_Credentials.csv"
xlsx_path_project = r"C:\Users\maddi\OneDrive\Desktop\DataI2I\Restaurant-Management-System\User_Credentials.xlsx"
xlsx_path_artifact = r"C:\Users\maddi\.gemini\antigravity-ide\brain\5fd5189b-279a-4de3-8518-8664366b318d\User_Credentials.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "User Credentials"

header_fill = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid") # Vibrant Orange header
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
cell_font = Font(name="Calibri", size=10)
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

with open(csv_path, 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        ws.append(row)
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if row_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = cell_font
                cell.alignment = Alignment(horizontal="left", vertical="center")

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save(xlsx_path_project)
wb.save(xlsx_path_artifact)
print(f"Excel file successfully generated at {xlsx_path_project} and {xlsx_path_artifact}")
